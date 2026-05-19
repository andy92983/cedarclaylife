#!/usr/bin/env python3
"""Generate branded Cedar & Clay documents (Prices, FMPlants$, elderberry directions)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
LOGO = ROOT / "Logos" / "Logo.png"
ICON = ROOT / "Logos" / "Favicon.png"

BRAND = {
    "name": "Cedar & Clay",
    "website": "cedarclaylife.com",
    "phone": "(920) 538-6638",
    "contact": "Samantha Bergstrom",
    "address": "W4187 Blackhawk Ave, Redgranite, WI 54970",
}

CEDAR = RGBColor(0x7A, 0x55, 0x38)
CLAY = RGBColor(0xB8, 0x6A, 0x50)
SAGE = RGBColor(0x63, 0x7A, 0x56)
BARK = RGBColor(0x2E, 0x22, 0x19)
MUTED = RGBColor(0x96, 0x6B, 0x45)


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_horizontal_rule(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B86A50")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def style_run(run, *, size: int = 11, bold: bool = False, italic: bool = False, color=BARK):
    run.font.name = "Georgia"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_brand_header(doc: Document) -> None:
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(3.2))
    else:
        p = doc.add_paragraph(BRAND["name"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.runs[0], size=18, bold=True, color=CEDAR)

    site = doc.add_paragraph(BRAND["website"])
    site.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(site.runs[0], size=9, color=SAGE)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_horizontal_rule(rule)


def add_product_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title.upper())
    style_run(run, size=16, bold=True, color=CEDAR)
    p.paragraph_format.space_after = Pt(6)


def add_labeled_line(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}: ")
    style_run(label_run, size=10, bold=True, color=CEDAR)
    text_run = p.add_run(text)
    style_run(text_run, size=10, color=BARK)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix:
        br = p.add_run(f"{bold_prefix}: ")
        style_run(br, size=10, bold=True, color=CEDAR)
    run = p.add_run(text)
    style_run(run, size=10, color=BARK)
    p.paragraph_format.space_after = Pt(4)


def add_contact_footer(doc: Document) -> None:
    doc.add_paragraph()
    for line in [BRAND["contact"], BRAND["phone"], BRAND["address"], BRAND["website"]]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(line), size=8, color=MUTED)


def build_prices_sheet() -> Document:
    doc = Document()
    add_brand_header(doc)
    add_product_title(doc, "Farmer's Market Prices")

    items = [
        "Elderberry syrup kits",
        "Elderberry syrup",
        "Vanilla",
        "Vanilla sugar",
        "Eggs",
        "Laundry Soap",
        "Fire starters",
        "Pink Himalayan Detox Bath",
        "Bath Soak",
        "Foot / Armpit Mask",
        "Foot soak",
        "Tallow Lotion",
        "Sugar Scrub (Peppermint)",
        "Lily of the Valley",
        "Comfrey",
        "Motherwort",
        "Daylilies",
        "Iris",
    ]

    table = doc.add_table(rows=len(items) + 1, cols=2)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "Product"
    headers[1].text = "Price"
    for cell in headers:
        set_cell_shading(cell, "E8DFD0")
        for p in cell.paragraphs:
            for run in p.runs:
                style_run(run, size=10, bold=True, color=CEDAR)

    for idx, item in enumerate(items, start=1):
        row = table.rows[idx].cells
        row[0].text = item
        row[1].text = ""
        for p in row[0].paragraphs:
            for run in p.runs:
                style_run(run, size=10, color=BARK)

    doc.add_paragraph()
    add_contact_footer(doc)
    return doc


def build_directions_sheet(title: str, blocks: list[dict]) -> Document:
    doc = Document()
    add_brand_header(doc)
    add_product_title(doc, title)

    for i, block in enumerate(blocks):
        if i > 0:
            doc.add_paragraph()
            rule = doc.add_paragraph()
            add_horizontal_rule(rule)
            doc.add_paragraph()

        add_product_title(doc, block["name"])
        if block.get("directions"):
            add_labeled_line(doc, "Directions", block["directions"])
        if block.get("ingredients"):
            add_labeled_line(doc, "Ingredients", block["ingredients"])
        if block.get("caution"):
            add_body(doc, block["caution"], bold_prefix="Caution")
        add_contact_footer(doc)

    return doc


def build_elderberry_directions() -> Document:
    return build_directions_sheet(
        "Elderberry Syrup Kit — Instructions",
        [
            {
                "name": "Stovetop Method",
                "directions": (
                    "Pour contents into a pan with 4 cups water. Bring to a boil, then reduce heat "
                    "and simmer 1 hour or until liquid is reduced by about half. Cool to room "
                    "temperature, strain, and add 1 cup honey. Store refrigerated up to 2 months. "
                    "Serving size: adults 1 tbsp, children 1 tsp."
                ),
                "ingredients": "Elderberries, rosehip, cinnamon, anise star, cloves.",
                "caution": (
                    "Do not consume raw elderberries. Honey is not safe for children under one year old. "
                    "Product produced in a private residence exempt from state licensing and inspection."
                ),
            },
            {
                "name": "Instant Pot Method",
                "directions": (
                    "Add contents and water to the pot. Seal the lid and cook on manual high pressure "
                    "for 8 minutes. Quick-release steam, mash berries, strain juice, and cool. "
                    "Stir in honey until dissolved. Refrigerate up to 2 months."
                ),
                "ingredients": "Elderberries, rosehip, cinnamon, anise star, cloves.",
                "caution": (
                    "Do not consume raw elderberries. Honey is not safe for children under one year old. "
                    "Product produced in a private residence exempt from state licensing and inspection."
                ),
            },
        ],
    )


def build_plants_sheet() -> Document:
    doc = Document()
    add_brand_header(doc)
    add_product_title(doc, "Garden Plants")

    plants = [
        "Lily of the Valley",
        "Comfrey",
        "Motherwort",
        "Daylilies",
        "Iris",
    ]

    for plant in plants:
        p = doc.add_paragraph()
        name = p.add_run(plant)
        style_run(name, size=12, bold=True, color=CEDAR)
        price = p.add_run("  —  Price: __________")
        style_run(price, size=11, color=CLAY)
        p.paragraph_format.space_after = Pt(10)

    doc.add_paragraph()
    note = doc.add_paragraph("Seasonal availability from our acreage in Redgranite, Wisconsin.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(note.runs[0], size=9, italic=True, color=MUTED)
    add_contact_footer(doc)
    return doc


def main() -> None:
    print("Generating branded documents...")
    build_prices_sheet().save(ROOT / "Prices.docx")
    print("  wrote Prices.docx")
    build_plants_sheet().save(ROOT / "FMPlants$.docx")
    print("  wrote FMPlants$.docx")
    build_elderberry_directions().save(ROOT / "Directions for Elder berries.docx")
    print("  wrote Directions for Elder berries.docx")
    print("Done.")


if __name__ == "__main__":
    main()
